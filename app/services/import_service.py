"""
Parse an uploaded CSV/XLSX file into a typed table definition and rows.

Type inference is deliberately conservative: a column is only given a
non-text type when every non-empty sampled value parses as that type.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

from openpyxl import load_workbook

from app.models.mcp_tools import IDENTIFIER_PATTERN

MAX_ROWS = 500_000
MAX_BYTES = 200 * 1024 * 1024
SAMPLE_ROWS = 1_000

_TRUE = {"true", "t", "yes", "y", "1"}
_FALSE = {"false", "f", "no", "n", "0"}
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y")
_TS_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%SZ")


class ImportError_(ValueError):
    """Raised for user-correctable problems with the uploaded file."""


@dataclass
class ColumnSpec:
    name: str
    type: str  # postgres type name
    source_header: str


@dataclass
class ImportPlan:
    columns: list[ColumnSpec]
    rows: list[list[Any]] = field(default_factory=list)
    skipped_rows: int = 0

    @property
    def row_count(self) -> int:
        return len(self.rows)


def parse_upload(filename: str, content: bytes) -> ImportPlan:
    """Detect the format from the filename, parse, infer types and coerce values."""
    if len(content) > MAX_BYTES:
        raise ImportError_(f"File exceeds {MAX_BYTES // (1024 * 1024)} MB")

    lower = filename.lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        headers, raw_rows = _read_csv(content)
    elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        headers, raw_rows = _read_xlsx(content)
    else:
        raise ImportError_("Unsupported file type; upload .csv or .xlsx")

    if not headers:
        raise ImportError_("The file has no header row")

    if len(raw_rows) > MAX_ROWS:
        raise ImportError_(f"File has {len(raw_rows)} rows; the limit is {MAX_ROWS}")

    names = sanitize_headers(headers)
    width = len(headers)
    types = [_infer_type(col_values) for col_values in _columns(raw_rows, width)]

    # Inference only looked at a sample; if a later row disagrees, demote
    # just that column to text and start over. Bounded by the column count.
    while True:
        try:
            rows, skipped = _coerce_rows(raw_rows, types, width)
            break
        except _ColumnMismatch as m:
            types[m.index] = "text"

    columns = [ColumnSpec(name=n, type=t, source_header=h) for n, t, h in zip(names, types, headers)]
    return ImportPlan(columns=columns, rows=rows, skipped_rows=skipped)


class _ColumnMismatch(Exception):
    def __init__(self, index: int) -> None:
        self.index = index


def _coerce_rows(raw_rows, types, width) -> tuple[list[list[Any]], int]:
    rows: list[list[Any]] = []
    skipped = 0
    for raw in raw_rows:
        if all(_is_empty(v) for v in raw):
            skipped += 1
            continue
        out = []
        for i, (v, t) in enumerate(zip(_pad(raw, width), types)):
            try:
                out.append(_coerce(v, t))
            except (ValueError, InvalidOperation, TypeError):
                raise _ColumnMismatch(i)
        rows.append(out)
    return rows, skipped


def default_table_name(filename: str) -> str:
    stem = re.sub(r"\.[^.]+$", "", filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1])
    return sanitize_headers([stem or "imported"])[0]


def sanitize_headers(headers: Iterable[str]) -> list[str]:
    """Turn arbitrary header text into unique, valid Postgres identifiers."""
    seen: dict[str, int] = {}
    out: list[str] = []
    for i, h in enumerate(headers):
        name = re.sub(r"[^a-z0-9_]+", "_", str(h or "").strip().lower()).strip("_")
        if not name:
            name = f"column_{i + 1}"
        if name[0].isdigit():
            name = f"c_{name}"
        name = name[:60]
        base = name
        while name in seen:
            seen[base] = seen.get(base, 1) + 1
            name = f"{base}_{seen[base]}"
        seen[name] = 1
        assert re.fullmatch(IDENTIFIER_PATTERN, name)
        out.append(name)
    return out


# ---------------------------------------------------------------- readers

def _read_csv(content: bytes) -> tuple[list[str], list[list[Any]]]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportError_("Could not decode the file as text")

    sample = text[:10_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _read_xlsx(content: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ImportError_(f"Could not read the workbook: {e}")
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        return [], []
    headers = ["" if h is None else str(h) for h in rows[0]]
    return headers, rows[1:]


# ---------------------------------------------------------------- inference

def _columns(rows: list[list[Any]], width: int) -> list[list[Any]]:
    sample = rows[:SAMPLE_ROWS]
    return [[_pad(r, width)[i] for r in sample] for i in range(width)]


def _pad(row: list[Any], width: int) -> list[Any]:
    return list(row[:width]) + [None] * (width - len(row))


def _is_empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _infer_type(values: list[Any]) -> str:
    present = [v for v in values if not _is_empty(v)]
    if not present:
        return "text"
    for candidate in ("bigint", "numeric", "boolean", "date", "timestamp"):
        if all(_parses_as(v, candidate) for v in present):
            return candidate
    return "text"


def _parses_as(v: Any, t: str) -> bool:
    try:
        _coerce(v, t)
        return True
    except (ValueError, InvalidOperation, TypeError):
        return False


def _naive_utc(dt: datetime) -> datetime:
    """asyncpg needs naive values for `timestamp`; aware inputs are converted to UTC."""
    if dt.tzinfo is not None:
        return dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


def _coerce(v: Any, t: str) -> Any:
    if _is_empty(v):
        return None
    if t == "text":
        return v if isinstance(v, str) else str(v)
    if t == "bigint":
        if isinstance(v, bool):
            raise ValueError
        if isinstance(v, int):
            return v
        if isinstance(v, float) and v.is_integer():
            return int(v)
        s = str(v).strip().replace(",", "")
        if not re.fullmatch(r"[+-]?\d+", s):
            raise ValueError
        n = int(s)
        if not -(2**63) <= n < 2**63:
            raise ValueError
        return n
    if t == "numeric":
        if isinstance(v, bool):
            raise ValueError
        if isinstance(v, (int, float, Decimal)):
            return Decimal(str(v))
        return Decimal(str(v).strip().replace(",", ""))
    if t == "boolean":
        if isinstance(v, bool):
            return v
        s = str(v).strip().lower()
        if s in _TRUE:
            return True
        if s in _FALSE:
            return False
        raise ValueError
    if t == "date":
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        s = str(v).strip()
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        raise ValueError
    if t == "timestamp":
        if isinstance(v, datetime):
            return _naive_utc(v)
        s = str(v).strip()
        for fmt in _TS_FORMATS:
            try:
                return _naive_utc(datetime.strptime(s, fmt))
            except ValueError:
                continue
        return _naive_utc(datetime.fromisoformat(s))
    raise ValueError(f"unknown type {t}")
