"""
Schema changes, dry runs, undo, storage limits, browsing and the SQL
editor, sharing and import/export, all through the HTTP API.
"""
import io

import pytest
from openpyxl import load_workbook

from app.core.config import get_settings

pytestmark = pytest.mark.asyncio


async def sql(client, user, database_id, query):
    return await client.post(f"/databases/{database_id}/sql", json={"sql": query}, headers=user.headers)


async def rows(client, user, database_id, query):
    r = await sql(client, user, database_id, query)
    assert r.status_code == 200, r.text
    return r.json()["rows"]


async def confirm(client, user, database_id, query):
    r = await client.post(f"/databases/{database_id}/sql/confirm", json={"sql": query}, headers=user.headers)
    assert r.status_code == 200, r.text
    return r.json()


async def undo(client, user, database_id):
    return await client.post(f"/databases/{database_id}/undo", headers=user.headers)


async def share(client, owner, database_id, user, level):
    r = await client.post(
        f"/databases/{database_id}/share", json={"email": user.email, "access_level": level}, headers=owner.headers
    )
    assert r.status_code == 204, r.text


# ------------------------------------------------------------ dry runs

async def test_writes_are_dry_run_before_confirmation(client, owner, seeded):
    r = await sql(client, owner, seeded, "DELETE FROM customers WHERE city = 'Pune'")
    assert r.status_code == 409
    detail = r.json()["detail"]
    assert detail["code"] == "confirmation_required"
    assert detail["rows_affected"] == 2
    assert detail["undo_available"] is True
    # The dry run was rolled back.
    assert await rows(client, owner, seeded, "SELECT count(*) FROM customers") == [[3]]


async def test_broken_writes_fail_before_confirmation(client, owner, seeded):
    r = await sql(client, owner, seeded, "UPDATE customers SET nope = 1")
    assert r.status_code == 400 and "nope" in r.text


async def test_truncate_preview_counts_discarded_rows(client, owner, seeded):
    r = await sql(client, owner, seeded, "TRUNCATE customers")
    assert r.json()["detail"]["rows_affected"] == 3


# ------------------------------------------------------------ schema changes

async def test_schema_changes_are_owner_only(client, owner, make_user, seeded):
    writer = await make_user()
    await share(client, owner, seeded, writer, "write")

    r = await sql(client, writer, seeded, "ALTER TABLE customers ADD COLUMN email text")
    assert r.status_code == 400 and "only allowed for the database owner" in r.text
    r = await sql(client, writer, seeded, "DROP TABLE customers")
    assert r.status_code == 400

    # Ordinary writes are still fine for them.
    assert (await sql(client, writer, seeded, "DELETE FROM customers WHERE id = 1")).status_code == 409

    assert (await sql(client, owner, seeded, "ALTER TABLE customers ADD COLUMN email text")).status_code == 409
    result = await confirm(client, owner, seeded, "ALTER TABLE customers ADD COLUMN email text")
    assert result["undo_available"] is True
    columns = (await client.get(f"/databases/{seeded}/tables/customers", headers=owner.headers)).json()["columns"]
    assert "email" in [c["name"] for c in columns]


async def test_ownership_and_schema_moves_are_rejected(client, owner, seeded):
    for statement in (
        "ALTER TABLE customers OWNER TO postgres",
        "ALTER TABLE customers SET SCHEMA public",
        "DROP SCHEMA public",
        "DROP FUNCTION f",
    ):
        r = await sql(client, owner, seeded, statement)
        assert r.status_code == 400, statement


# ------------------------------------------------------------ undo

async def test_undo_delete_and_update(client, owner, seeded):
    await confirm(client, owner, seeded, "DELETE FROM customers WHERE city = 'Pune'")
    assert await rows(client, owner, seeded, "SELECT count(*) FROM customers") == [[1]]

    status = (await client.get(f"/databases/{seeded}/undo", headers=owner.headers)).json()
    assert status["available"] and status["sql"].startswith("DELETE") and status["by_you"]

    r = await undo(client, owner, seeded)
    assert r.status_code == 200 and r.json()["undone_sql"].startswith("DELETE")
    assert await rows(client, owner, seeded, "SELECT name FROM customers ORDER BY id") == [["Ann"], ["Bob"], ["Cy"]]

    # One level only.
    assert (await undo(client, owner, seeded)).status_code == 404

    await confirm(client, owner, seeded, "UPDATE customers SET orders = 99")
    await undo(client, owner, seeded)
    assert await rows(client, owner, seeded, "SELECT orders FROM customers ORDER BY id") == [[3], [0], [1]]


async def test_undo_drop_table_restores_data_keys_and_serial(client, owner, database):
    await confirm(client, owner, database, "CREATE TABLE items (id serial PRIMARY KEY, label text NOT NULL)")
    await confirm(client, owner, database, "INSERT INTO items (label) VALUES ('a'), ('b')")
    await confirm(client, owner, database, "DROP TABLE items")
    assert (await client.get(f"/databases/{database}/tables", headers=owner.headers)).json()["tables"] == []

    assert (await undo(client, owner, database)).status_code == 200
    assert await rows(client, owner, database, "SELECT id, label FROM items ORDER BY id") == [[1, "a"], [2, "b"]]
    columns = (await client.get(f"/databases/{database}/tables/items", headers=owner.headers)).json()["columns"]
    assert [c["primary_key"] for c in columns] == [True, False]
    # The serial default came back and continues after the existing ids.
    await confirm(client, owner, database, "INSERT INTO items (label) VALUES ('c')")
    assert await rows(client, owner, database, "SELECT max(id) FROM items") == [[3]]


async def test_undo_alter_create_and_truncate(client, owner, seeded):
    await confirm(client, owner, seeded, "ALTER TABLE customers RENAME COLUMN city TO town")
    await undo(client, owner, seeded)
    assert await rows(client, owner, seeded, "SELECT city FROM customers WHERE id = 1") == [["Pune"]]

    await confirm(client, owner, seeded, "CREATE TABLE scratch (x int)")
    await undo(client, owner, seeded)
    assert "scratch" not in (await client.get(f"/databases/{seeded}/tables", headers=owner.headers)).json()["tables"]

    await confirm(client, owner, seeded, "CREATE VIEW pune AS SELECT * FROM customers WHERE city = 'Pune'")
    await confirm(client, owner, seeded, "TRUNCATE customers")
    await undo(client, owner, seeded)
    assert await rows(client, owner, seeded, "SELECT count(*) FROM pune") == [[2]]


async def test_undo_restores_dependent_views_after_drop_cascade(client, owner, seeded):
    await confirm(client, owner, seeded, "CREATE VIEW pune AS SELECT * FROM customers WHERE city = 'Pune'")
    await confirm(client, owner, seeded, "DROP TABLE customers CASCADE")
    await undo(client, owner, seeded)
    assert await rows(client, owner, seeded, "SELECT name FROM pune ORDER BY id") == [["Ann"], ["Cy"]]


async def test_undo_unavailable_when_referenced_by_foreign_key(client, owner, database):
    await confirm(client, owner, database, "CREATE TABLE parent (id int PRIMARY KEY)")
    await confirm(client, owner, database, "CREATE TABLE child (parent_id int REFERENCES parent(id))")
    r = await sql(client, owner, database, "DELETE FROM parent")
    assert r.json()["detail"]["undo_available"] is False
    assert "foreign key" in r.json()["detail"]["undo_unavailable_reason"]

    result = await confirm(client, owner, database, "DELETE FROM parent")
    assert result["undo_available"] is False
    assert (await undo(client, owner, database)).status_code == 409


async def test_import_invalidates_undo_and_readers_cannot_undo(client, owner, make_user, seeded):
    reader = await make_user()
    await share(client, owner, seeded, reader, "read")
    await confirm(client, owner, seeded, "DELETE FROM customers WHERE id = 1")
    assert (await undo(client, reader, seeded)).status_code == 403

    r = await client.post(
        f"/databases/{seeded}/tables/import",
        files={"file": ("extra.csv", b"a\n1\n")},
        headers=owner.headers,
    )
    assert r.status_code == 200, r.text
    assert (await undo(client, owner, seeded)).status_code == 404
    snapshots = await rows(client, owner, seeded, "SELECT count(*) FROM pg_tables WHERE tablename LIKE '\\_undo\\_%'")
    assert snapshots == [[0]]


async def test_snapshots_are_hidden_from_table_lists(client, owner, seeded):
    await confirm(client, owner, seeded, "UPDATE customers SET orders = 0")
    tables = (await client.get(f"/databases/{seeded}/tables", headers=owner.headers)).json()["tables"]
    assert tables == ["customers"]


# ------------------------------------------------------------ storage limit

async def test_storage_limit_rolls_back_growth_but_allows_deletes(client, owner, seeded, monkeypatch):
    monkeypatch.setattr(get_settings(), "database_max_mb", 1)
    big = "INSERT INTO customers (id, name, city, orders) SELECT g, repeat('x', 500), 'Z', 0 FROM generate_series(10, 5000) g"
    r = await client.post(f"/databases/{seeded}/sql/confirm", json={"sql": big}, headers=owner.headers)
    assert r.status_code == 413 and "storage limit" in r.text
    assert await rows(client, owner, seeded, "SELECT count(*) FROM customers") == [[3]]

    await confirm(client, owner, seeded, "DELETE FROM customers WHERE id = 1")

    usage = (await client.get(f"/databases/{seeded}/usage", headers=owner.headers)).json()
    assert usage["limit_bytes"] == 1024 * 1024 and 0 < usage["size_bytes"] < usage["limit_bytes"]


async def test_storage_limit_applies_to_imports(client, owner, database, monkeypatch):
    monkeypatch.setattr(get_settings(), "database_max_mb", 1)
    csv = "id,text\n" + "".join(f"{i},{'y' * 400}\n" for i in range(5000))
    r = await client.post(
        f"/databases/{database}/tables/import", files={"file": ("big.csv", csv.encode())}, headers=owner.headers
    )
    assert r.status_code == 413
    assert (await client.get(f"/databases/{database}/tables", headers=owner.headers)).json()["tables"] == []


# ------------------------------------------------------------ browsing and SQL editor

async def test_browse_rows_paginates_and_sorts(client, owner, seeded):
    r = await client.get(
        f"/databases/{seeded}/tables/customers/rows",
        params={"limit": 2, "offset": 1, "order_by": "orders", "descending": True},
        headers=owner.headers,
    )
    body = r.json()
    assert body["total"] == 3 and body["columns"][:2] == ["id", "name"]
    assert [row[1] for row in body["rows"]] == ["Cy", "Bob"]

    r = await client.get(
        f"/databases/{seeded}/tables/customers/rows", params={"order_by": "nope"}, headers=owner.headers
    )
    assert r.status_code == 400
    r = await client.get(f"/databases/{seeded}/tables/missing/rows", headers=owner.headers)
    assert r.status_code == 404


async def test_sql_editor_caps_rows(client, owner, database):
    r = await sql(client, owner, database, "SELECT g FROM generate_series(1, 1005) g;")
    body = r.json()
    assert body["truncated"] is True and len(body["rows"]) == 1000
    r = await sql(client, owner, database, "SELECT 1 AS one")
    assert r.json() == {"columns": ["one"], "rows": [[1]], "truncated": False}


# ------------------------------------------------------------ sharing

async def test_sharing_levels_and_revocation(client, owner, make_user, seeded):
    writer = await make_user()
    await share(client, owner, seeded, writer, "write")

    # Writers can write but not share or delete.
    assert (await sql(client, writer, seeded, "INSERT INTO customers VALUES (9, 'Z', 'Goa', 0)")).status_code == 409
    other = await make_user()
    r = await client.post(
        f"/databases/{seeded}/share", json={"email": other.email, "access_level": "read"}, headers=writer.headers
    )
    assert r.status_code == 403
    assert (await client.delete(f"/databases/{seeded}", headers=writer.headers)).status_code == 403

    grants = (await client.get(f"/databases/{seeded}/permissions", headers=owner.headers)).json()
    assert [(g["email"], g["access_level"]) for g in grants] == [(writer.email, "write")]

    writer_id = await writer.me()
    assert (await client.delete(f"/databases/{seeded}/permissions/{writer_id}", headers=owner.headers)).status_code == 204
    assert (await client.get(f"/databases/{seeded}/tables", headers=writer.headers)).status_code == 403


async def test_cannot_share_with_unverified_address(client, owner, make_user, seeded):
    unverified = await make_user(verified=False)
    r = await client.post(
        f"/databases/{seeded}/share", json={"email": unverified.email, "access_level": "read"}, headers=owner.headers
    )
    assert r.status_code == 409 and "verified" in r.text


# ------------------------------------------------------------ import / export

async def test_import_dry_run_append_and_type_inference(client, owner, database):
    csv = b"id,price,joined,active\n1,9.5,2024-01-02,yes\n2,10,2024-02-03,no\n"
    r = await client.post(
        f"/databases/{database}/tables/import",
        files={"file": ("Prices List.csv", csv)}, data={"dry_run": "true"}, headers=owner.headers,
    )
    body = r.json()
    assert body["dry_run"] and body["table_name"] == "prices_list"
    assert [c["type"] for c in body["columns"]] == ["bigint", "numeric", "date", "boolean"]
    assert (await client.get(f"/databases/{database}/tables", headers=owner.headers)).json()["tables"] == []

    await client.post(f"/databases/{database}/tables/import", files={"file": ("Prices List.csv", csv)}, headers=owner.headers)
    r = await client.post(
        f"/databases/{database}/tables/import",
        files={"file": ("more.csv", b"id,price,joined,active\n3,1,2024-03-04,true\n")},
        data={"table_name": "prices_list", "mode": "append"},
        headers=owner.headers,
    )
    assert r.status_code == 200
    assert await rows(client, owner, database, "SELECT count(*) FROM prices_list WHERE active") == [[2]]


async def test_export_table_as_csv_and_xlsx(client, owner, seeded):
    r = await client.get(f"/databases/{seeded}/tables/customers/export", params={"format": "csv"}, headers=owner.headers)
    assert r.status_code == 200
    lines = r.text.strip().splitlines()
    assert lines[0] == "id,name,city,orders" and len(lines) == 4

    r = await client.post(
        f"/databases/{seeded}/export",
        json={"sql": "SELECT name, orders FROM customers ORDER BY id", "format": "xlsx"},
        headers=owner.headers,
    )
    sheet = load_workbook(io.BytesIO(r.content)).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [["name", "orders"], ["Ann", 3], ["Bob", 0], ["Cy", 1]]


async def test_undo_refuses_when_write_reaches_unsnapshotted_tables(client, owner, seeded):
    # Inserting through an updatable view changes its base table, which the
    # SQL alone doesn't name; Postgres's transaction stats catch it.
    await confirm(client, owner, seeded, "CREATE TABLE notes (body text)")
    await confirm(client, owner, seeded, "CREATE VIEW note_view AS SELECT * FROM notes")
    result = await confirm(
        client, owner, seeded,
        "WITH x AS (INSERT INTO note_view VALUES ('hi') RETURNING 1) UPDATE customers SET orders = 0",
    )
    assert result["undo_available"] is False
    status = (await client.get(f"/databases/{seeded}/undo", headers=owner.headers)).json()
    assert status["unavailable_reason"] == "it also changed notes"
